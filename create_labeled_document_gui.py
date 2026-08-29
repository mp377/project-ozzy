"""
create_labeled_document_gui.py

Desktop interface for creating Microsoft Office documents with a Purview
sensitivity label applied at creation time.

Backed by a persistent ("warm") PowerShell session. Measured on this
environment: a cold PowerShell process costs ~92 seconds because the Purview
client re-bootstraps its engine every time. Holding one session open reduces
each labeling call to roughly a second.

The window stays open so the session stays warm. Create as many documents as
you like; only the first pays the startup cost.

REQUIREMENTS
------------
    pip install python-docx openpyxl python-pptx

    Purview Information Protection client installed, and authenticated:
        Set-Authentication

RUN
---
    python create_labeled_document_gui.py
"""

import os
import queue
import subprocess
import sys
import threading
import time
import uuid

try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
except ImportError:
    print("tkinter is not available. It ships with Python on Windows.")
    sys.exit(3)


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------

# Tenant-specific. Retrieve with: Get-Label | Select-Object DisplayName, Guid
LABELS = {
    "Public":            "994c2e99-9b3f-459d-8e3e-36de4c4f4f91",
    "Limited":           "7430563a-23b6-4772-b5b1-9317d899d4b3",
    "Restricted":        "7cd5dfcb-9dc7-4f57-985b-cabf186e9678",
    "Highly Restricted": "82c90564-5d3b-485b-933a-93e69a6e60ba",
}

LABEL_COLORS = {
    "Public":            ("#e8f4fd", "#005a9e"),
    "Limited":           ("#e6f4ea", "#1e7e34"),
    "Restricted":        ("#fff8e1", "#856404"),
    "Highly Restricted": ("#fce8e6", "#c62828"),
}

DOC_TYPES = {
    "Word document (.docx)":     "docx",
    "Excel workbook (.xlsx)":    "xlsx",
    "PowerPoint deck (.pptx)":   "pptx",
}

BROWSE_ENTRY = "Browse for another folder..."

COMMAND_TIMEOUT = 300
WARMUP_TIMEOUT = 300

# Palette
NAVY = "#1E2761"
SLATE = "#3D4A6B"
MUTED = "#6B7794"
LIGHTBG = "#F7F9FC"
BORDER = "#DDE4F0"


# --------------------------------------------------------------------------
# Candidate save locations
# --------------------------------------------------------------------------

def candidate_locations():
    """
    Build the list of folders offered in the dropdown.
    Only folders that actually exist are included, so the user cannot
    select an invalid destination.
    """
    home = os.path.expanduser("~")
    candidates = [
        ("Label Test folder",      os.path.join(home, "LabelTest")),
        ("Documents",              os.path.join(home, "Documents")),
        ("Desktop",                os.path.join(home, "Desktop")),
    ]

    # Any OneDrive roots present on this machine
    try:
        for entry in sorted(os.listdir(home)):
            full = os.path.join(home, entry)
            if entry.lower().startswith("onedrive") and os.path.isdir(full):
                docs = os.path.join(full, "Documents")
                if os.path.isdir(docs):
                    candidates.append((f"{entry} \u2192 Documents", docs))
                candidates.append((entry, full))
    except OSError:
        pass

    seen = set()
    result = []
    for name, path in candidates:
        norm = os.path.normpath(path)
        if os.path.isdir(norm) and norm.lower() not in seen:
            seen.add(norm.lower())
            result.append((name, norm))

    if not result:
        result.append(("Home folder", os.path.normpath(home)))

    return result


# --------------------------------------------------------------------------
# Warm PowerShell session
# --------------------------------------------------------------------------

class WarmSession:
    """One long-lived PowerShell process with the Purview engine bootstrapped."""

    def __init__(self, on_status=None):
        self.proc = None
        self._out_q = queue.Queue()
        self._started = False
        self._lock = threading.Lock()
        self.on_status = on_status or (lambda msg: None)

    # -- internals --------------------------------------------------------

    def _drain_stdout(self):
        try:
            for line in iter(self.proc.stdout.readline, ""):
                self._out_q.put(line.rstrip("\r\n"))
        except Exception:
            pass

    def _read_until(self, sentinel, timeout):
        lines = []
        deadline = time.time() + timeout
        while True:
            remaining = deadline - time.time()
            if remaining <= 0:
                raise TimeoutError(f"No response within {timeout}s")
            try:
                line = self._out_q.get(timeout=min(remaining, 1.0))
            except queue.Empty:
                if self.proc.poll() is not None:
                    raise RuntimeError("PowerShell session ended unexpectedly")
                continue
            if sentinel in line:
                return lines
            lines.append(line)

    # -- lifecycle --------------------------------------------------------

    def start(self):
        self.on_status("Starting PowerShell...")
        try:
            self.proc = subprocess.Popen(
                ["powershell", "-NoProfile", "-NoLogo", "-NonInteractive", "-Command", "-"],
                stdin=subprocess.PIPE,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                bufsize=1,
            )
        except FileNotFoundError:
            raise RuntimeError("PowerShell not found. This tool requires Windows.")

        threading.Thread(target=self._drain_stdout, daemon=True).start()
        self._started = True

        self.on_status("Loading Purview module...")
        self.run("Import-Module PurviewInformationProtection -ErrorAction SilentlyContinue",
                 timeout=60)

        self.on_status("Warming the labeling engine...")
        t0 = time.perf_counter()
        warned = self._warm_engine()
        elapsed = time.perf_counter() - t0

        if warned:
            self.on_status(f"Ready in {elapsed:.1f}s  ({warned})")
        else:
            self.on_status(f"Ready in {elapsed:.1f}s")
        return elapsed

    def _warm_engine(self):
        """
        Force a full bootstrap by labeling a throwaway file.
        Returns a warning string if the warmup did not complete cleanly.
        """
        import tempfile
        warm_path = os.path.join(tempfile.gettempdir(), f"_warmup_{uuid.uuid4().hex}.docx")

        try:
            create_word(warm_path, title="warmup")
        except Exception as e:
            return f"warmup skipped: {e}"

        try:
            safe = warm_path.replace("'", "''")
            out = self.run(
                f"try {{ Set-FileLabel -Path '{safe}' -LabelId '{LABELS['Public']}' "
                f"-ErrorAction Stop | Out-Null; Write-Output 'WARM_OK' }} "
                f"catch {{ Write-Output ('WARM_FAIL: ' + $_.Exception.Message) }}",
                timeout=WARMUP_TIMEOUT,
            )
            text = " ".join(out)
            if "WARM_OK" not in text:
                return "engine may not be fully warm"
        except Exception as e:
            return f"warmup error: {e}"
        finally:
            try:
                os.remove(warm_path)
            except OSError:
                pass
        return None

    def run(self, command, timeout=COMMAND_TIMEOUT):
        if not self._started or self.proc.poll() is not None:
            raise RuntimeError("Session is not running")
        with self._lock:
            sentinel = f"__DONE_{uuid.uuid4().hex}__"
            self.proc.stdin.write(f"{command}\nWrite-Output '{sentinel}'\n")
            self.proc.stdin.flush()
            return self._read_until(sentinel, timeout)

    def stop(self):
        if self.proc and self.proc.poll() is None:
            try:
                self.proc.stdin.write("exit\n")
                self.proc.stdin.flush()
                self.proc.wait(timeout=5)
            except Exception:
                try:
                    self.proc.kill()
                except Exception:
                    pass
        self._started = False

    @property
    def alive(self):
        return self._started and self.proc is not None and self.proc.poll() is None

    # -- operations -------------------------------------------------------

    def label_file(self, path, label_name):
        label_id = LABELS.get(label_name)
        if not label_id:
            return False, f"Unknown label: {label_name}"
        safe = path.replace("'", "''")
        cmd = (
            f"try {{ Set-FileLabel -Path '{safe}' -LabelId '{label_id}' "
            f"-ErrorAction Stop | Out-Null; Write-Output 'OK' }} "
            f"catch {{ Write-Output ('FAIL: ' + $_.Exception.Message) }}"
        )
        try:
            out = self.run(cmd)
        except Exception as e:
            return False, str(e)
        text = " ".join(out).strip()
        if "OK" in text and "FAIL" not in text:
            return True, "Label applied"
        return False, text or "No response from Set-FileLabel"

    def get_label(self, path):
        safe = path.replace("'", "''")
        cmd = (
            f"$s = Get-FileStatus -Path '{safe}' -ErrorAction SilentlyContinue; "
            f"if ($s -and $s.IsLabeled) {{ Write-Output $s.MainLabelName }} "
            f"else {{ Write-Output '' }}"
        )
        try:
            out = self.run(cmd, timeout=120)
        except Exception:
            return None
        text = " ".join(l for l in out if l.strip()).strip()
        return text or None


# --------------------------------------------------------------------------
# Document creation
# --------------------------------------------------------------------------

def create_word(path, title=None, content=None):
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    if content:
        for line in content.split("\n"):
            doc.add_paragraph(line)
    doc.save(path)


def create_excel(path, title=None, content=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if title:
        ws.title = title[:31]
        ws["A1"] = title
    if content:
        for i, line in enumerate(content.split("\n"), start=3):
            ws.cell(row=i, column=1, value=line)
    wb.save(path)


def create_powerpoint(path, title=None, content=None):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title or "Untitled"
    if content:
        slide.placeholders[1].text = content
    prs.save(path)


CREATORS = {"docx": create_word, "xlsx": create_excel, "pptx": create_powerpoint}


def create_document(path, title=None, content=None):
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext not in CREATORS:
        raise ValueError(f"Unsupported type: .{ext}")
    parent = os.path.dirname(os.path.abspath(path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)
    CREATORS[ext](path, title, content)


def build_path(folder, name, ext):
    """
    Assemble a full path.

    Only strips a trailing extension if it is one we recognise, so a name
    like "Q3.Report" keeps its dot instead of losing everything after it.
    """
    name = name.strip()
    stem, typed_ext = os.path.splitext(name)
    if typed_ext.lower().lstrip(".") in CREATORS:
        name = stem
    return os.path.normpath(os.path.join(folder, f"{name}.{ext}"))


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------

class App:
    def __init__(self, root):
        self.root = root
        self.session = None
        self.session_ready = False
        self.busy = False
        self.ui_q = queue.Queue()

        self.locations = candidate_locations()
        self.custom_folder = None

        root.title("Labeled Document Creator")
        root.geometry("640x660")
        root.configure(bg=LIGHTBG)
        root.protocol("WM_DELETE_WINDOW", self.on_close)

        self._build_ui()
        self._poll_queue()
        self._start_session()

    # -- layout -----------------------------------------------------------

    def _build_ui(self):
        header = tk.Frame(self.root, bg=NAVY, height=64)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header, text="Labeled Document Creator", bg=NAVY, fg="white",
                 font=("Segoe UI", 15, "bold")).pack(side="left", padx=18)

        self.status_dot = tk.Label(header, text="\u25cf", bg=NAVY, fg="#c9a227",
                                   font=("Segoe UI", 13))
        self.status_dot.pack(side="right", padx=(0, 6))
        self.status_lbl = tk.Label(header, text="Starting...", bg=NAVY, fg="#CADCFC",
                                   font=("Segoe UI", 9))
        self.status_lbl.pack(side="right")

        body = tk.Frame(self.root, bg=LIGHTBG)
        body.pack(fill="both", expand=True, padx=22, pady=16)
        body.columnconfigure(1, weight=1)

        def field_label(text, row):
            tk.Label(body, text=text, bg=LIGHTBG, fg=SLATE,
                     font=("Segoe UI", 10, "bold")).grid(
                row=row, column=0, sticky="w", pady=(10, 2))

        # File name
        field_label("File name", 0)
        self.name_var = tk.StringVar()
        self.name_entry = ttk.Entry(body, textvariable=self.name_var,
                                    font=("Segoe UI", 10))
        self.name_entry.grid(row=1, column=0, columnspan=2, sticky="ew")

        # Document type
        field_label("Document type", 2)
        self.type_var = tk.StringVar(value=list(DOC_TYPES.keys())[0])
        self.type_combo = ttk.Combobox(body, textvariable=self.type_var,
                                       values=list(DOC_TYPES.keys()),
                                       state="readonly", font=("Segoe UI", 10))
        self.type_combo.grid(row=3, column=0, columnspan=2, sticky="ew")

        # Location
        field_label("Save to", 4)
        self.loc_var = tk.StringVar()
        self.loc_combo = ttk.Combobox(body, textvariable=self.loc_var,
                                      state="readonly", font=("Segoe UI", 10))
        self.loc_combo.grid(row=5, column=0, columnspan=2, sticky="ew")
        self.loc_combo.bind("<<ComboboxSelected>>", self.on_location_change)
        self._refresh_locations()

        self.path_preview = tk.Label(body, text="", bg=LIGHTBG, fg=MUTED,
                                     font=("Consolas", 8), anchor="w")
        self.path_preview.grid(row=6, column=0, columnspan=2, sticky="ew", pady=(3, 0))

        # Label
        field_label("Sensitivity label", 7)
        self.label_var = tk.StringVar(value="Limited")
        self.label_combo = ttk.Combobox(body, textvariable=self.label_var,
                                        values=list(LABELS.keys()),
                                        state="readonly", font=("Segoe UI", 10))
        self.label_combo.grid(row=8, column=0, columnspan=2, sticky="ew")
        self.label_combo.bind("<<ComboboxSelected>>", lambda e: self.on_label_change())

        self.label_swatch = tk.Label(body, text="", font=("Segoe UI", 10, "bold"),
                                     pady=7)
        self.label_swatch.grid(row=9, column=0, columnspan=2, sticky="ew", pady=(6, 0))

        # Heading
        field_label("Heading inside the document (optional)", 10)
        self.title_var = tk.StringVar()
        ttk.Entry(body, textvariable=self.title_var,
                  font=("Segoe UI", 10)).grid(row=11, column=0, columnspan=2, sticky="ew")

        # Open after
        self.open_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(body, text="Open the document after labeling",
                        variable=self.open_var).grid(
            row=12, column=0, columnspan=2, sticky="w", pady=(12, 0))

        # Create button
        self.create_btn = tk.Button(body, text="Create and label",
                                    command=self.on_create,
                                    bg=NAVY, fg="white", relief="flat",
                                    font=("Segoe UI", 11, "bold"),
                                    pady=9, cursor="hand2",
                                    activebackground="#005a9e",
                                    activeforeground="white")
        self.create_btn.grid(row=13, column=0, columnspan=2, sticky="ew", pady=(16, 0))
        self.create_btn.config(state="disabled")

        # Activity log
        tk.Label(body, text="Activity", bg=LIGHTBG, fg=SLATE,
                 font=("Segoe UI", 10, "bold")).grid(
            row=14, column=0, sticky="w", pady=(16, 2))

        log_frame = tk.Frame(body, bg="white", highlightbackground=BORDER,
                             highlightthickness=1)
        log_frame.grid(row=15, column=0, columnspan=2, sticky="nsew")
        body.rowconfigure(15, weight=1)

        self.log = tk.Text(log_frame, height=8, bg="white", fg=SLATE, relief="flat",
                           font=("Consolas", 9), wrap="word", state="disabled",
                           padx=8, pady=6)
        self.log.pack(side="left", fill="both", expand=True)
        sb = ttk.Scrollbar(log_frame, command=self.log.yview)
        sb.pack(side="right", fill="y")
        self.log.config(yscrollcommand=sb.set)

        self.name_var.trace_add("write", lambda *a: self.update_preview())
        self.type_var.trace_add("write", lambda *a: self.update_preview())
        self.on_label_change()
        self.update_preview()
        self.name_entry.focus_set()

    # -- helpers ----------------------------------------------------------

    def _refresh_locations(self, select=None):
        names = [f"{n}   ({p})" for n, p in self.locations]
        if self.custom_folder:
            names.append(f"Chosen folder   ({self.custom_folder})")
        names.append(BROWSE_ENTRY)
        self.loc_combo["values"] = names
        if select is not None:
            self.loc_var.set(select)
        elif not self.loc_var.get() or self.loc_var.get() == BROWSE_ENTRY:
            self.loc_var.set(names[0])

    def selected_folder(self):
        choice = self.loc_var.get()
        if choice == BROWSE_ENTRY:
            return None
        if self.custom_folder and choice.startswith("Chosen folder"):
            return self.custom_folder
        idx = self.loc_combo["values"].index(choice) if choice in self.loc_combo["values"] else 0
        if idx < len(self.locations):
            return self.locations[idx][1]
        return None

    def selected_ext(self):
        return DOC_TYPES.get(self.type_var.get(), "docx")

    def on_location_change(self, _event=None):
        if self.loc_var.get() == BROWSE_ENTRY:
            chosen = filedialog.askdirectory(title="Choose a folder")
            if chosen:
                self.custom_folder = os.path.normpath(chosen)
                self._refresh_locations()
                self._refresh_locations(
                    select=f"Chosen folder   ({self.custom_folder})")
            else:
                self._refresh_locations(select=self.loc_combo["values"][0])
        self.update_preview()

    def on_label_change(self):
        name = self.label_var.get()
        bg, fg = LABEL_COLORS.get(name, ("#eeeeee", "#333333"))
        self.label_swatch.config(text=name, bg=bg, fg=fg)

    def update_preview(self):
        folder = self.selected_folder()
        name = self.name_var.get().strip() or "(file name)"
        if folder:
            self.path_preview.config(
                text=build_path(folder, name, self.selected_ext()))
        else:
            self.path_preview.config(text="")

    def log_line(self, text, tag=None):
        self.log.config(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.config(state="disabled")

    def set_status(self, text, ready=False, error=False):
        self.status_lbl.config(text=text)
        self.status_dot.config(fg="#5cb85c" if ready else
                               ("#d9534f" if error else "#c9a227"))

    # -- threading --------------------------------------------------------

    def _poll_queue(self):
        try:
            while True:
                fn = self.ui_q.get_nowait()
                fn()
        except queue.Empty:
            pass
        self.root.after(80, self._poll_queue)

    def ui(self, fn):
        self.ui_q.put(fn)

    def _start_session(self):
        self.log_line("Starting the labeling session.")
        self.log_line("The first engine load can take up to two minutes.")

        def work():
            self.session = WarmSession(
                on_status=lambda m: self.ui(lambda: self.set_status(m)))
            try:
                elapsed = self.session.start()
            except Exception as e:
                self.ui(lambda: self.set_status("Session failed", error=True))
                self.ui(lambda: self.log_line(f"[ERROR] {e}"))
                self.ui(lambda: self.log_line(
                    "Check that the Purview client is installed and that "
                    "Set-Authentication has been run."))
                return
            self.session_ready = True
            self.ui(lambda: self.set_status("Session ready", ready=True))
            self.ui(lambda: self.log_line(
                f"Engine ready in {elapsed:.1f}s. Labels will now apply in about a second."))
            self.ui(lambda: self.create_btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    # -- actions ----------------------------------------------------------

    def on_create(self):
        if self.busy:
            return

        name = self.name_var.get().strip()
        if not name:
            messagebox.showwarning("Missing name", "Enter a file name.")
            self.name_entry.focus_set()
            return

        folder = self.selected_folder()
        if not folder:
            messagebox.showwarning("No folder", "Choose where to save the document.")
            return
        if not os.path.isdir(folder):
            messagebox.showerror("Folder not found", f"This folder does not exist:\n{folder}")
            return

        ext = self.selected_ext()
        path = build_path(folder, name, ext)
        label = self.label_var.get()
        title = self.title_var.get().strip() or None

        if os.path.exists(path):
            if not messagebox.askyesno("File exists",
                                       f"{os.path.basename(path)} already exists.\n\nOverwrite it?"):
                return

        self.busy = True
        self.create_btn.config(state="disabled", text="Working...")
        self.log_line("")
        self.log_line(f"Creating {os.path.basename(path)}")

        def work():
            try:
                create_document(path, title)
            except Exception as e:
                self.ui(lambda: self.log_line(f"[ERROR] Could not create it: {e}"))
                self.ui(self._done)
                return

            self.ui(lambda: self.log_line(f"  Saved to {folder}"))
            self.ui(lambda: self.log_line(f"  Applying '{label}'..."))

            t0 = time.perf_counter()
            ok, msg = self.session.label_file(os.path.abspath(path), label)
            elapsed = time.perf_counter() - t0

            if not ok:
                self.ui(lambda: self.log_line(f"  [ERROR] Labeling failed: {msg}"))
                self.ui(lambda: self.log_line("  The file exists but is unlabeled."))
                self.ui(self._done)
                return

            self.ui(lambda: self.log_line(f"  Labeled in {elapsed:.1f}s"))

            actual = self.session.get_label(os.path.abspath(path))
            if actual:
                verdict = "verified" if actual == label else f"MISMATCH (found '{actual}')"
                self.ui(lambda: self.log_line(f"  {verdict}"))
            else:
                self.ui(lambda: self.log_line("  [WARN] Could not read the label back."))

            if self.open_var.get():
                try:
                    os.startfile(os.path.abspath(path))
                    self.ui(lambda: self.log_line("  Opened in Office."))
                except Exception as e:
                    self.ui(lambda: self.log_line(f"  [WARN] Could not open it: {e}"))

            self.ui(self._done)

        threading.Thread(target=work, daemon=True).start()

    def _done(self):
        self.busy = False
        self.create_btn.config(state="normal", text="Create and label")

    def on_close(self):
        if self.busy:
            if not messagebox.askyesno("Still working",
                                       "A document is still being labeled. Close anyway?"):
                return
        if self.session:
            self.session.stop()
        self.root.destroy()


def main():
    root = tk.Tk()
    try:
        ttk.Style().theme_use("vista")
    except tk.TclError:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
